import json
import random
import io
import os
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
import streamlit as st

st.set_page_config(page_title="Convocazioni e Distinta Agape", page_icon="⚽", layout="centered")

@st.cache_data
def carica_rosa():
    if os.path.exists("giocatori.json"):
        try:
            with open("giocatori.json", "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            st.error(f"Errore nella lettura di giocatori.json: {e}")
            return []
    else:
        st.error("File 'giocatori.json' non trovato! Assicurati che sia stato caricato nel repository GitHub.")
        return []

ROSA = carica_rosa()
MIA_SQUADRA = "AGAPE UMBERTIDE PIERANTONIO"

def genera_convocazione_random(rosa):
    avanzati = [g for g in rosa if g["Livello"] == "Avanzato"]
    intermedi = [g for g in rosa if g["Livello"] == "Intermedio"]
    base = [g for g in rosa if g["Livello"] == "Base"]
    
    random.shuffle(avanzati)
    random.shuffle(intermedi)
    random.shuffle(base)
    
    squadra_rossa = avanzati[:4] + intermedi[:4] + base[:4]
    squadra_blu = avanzati[4:] + intermedi[4:] + base[4:]
    
    return squadra_rossa, squadra_blu

def formatta_giocatori(giocatori):
    if not giocatori:
        return "Nessun giocatore selezionato."
    return "\n".join([f"- {g['Cognome']} {g['Nome']}" for g in giocatori])

def crea_distinta_excel(
    squadra_rossa, squadra_blu, genera_distinta_r, genera_distinta_b,
    delegazione, stagione, categoria, anno, modalita, luogo, num_gare,
    data_r, avversario_rosso, in_casa_rosso, tecnico_r, tess_tec_r, dirigente_r, tess_dir_r, guardalinee_r, tess_gua_r,
    data_b, avversario_blu, in_casa_blu, tecnico_b, tess_tec_b, dirigente_b, tess_dir_b, guardalinee_b, tess_gua_b
):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    thin = Side(style='thin', color='000000')
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    f_title = Font(name="Arial", size=9, bold=True)
    f_sub = Font(name="Arial", size=8, bold=True)
    f_norm = Font(name="Arial", size=8)
    f_small = Font(name="Arial", size=7)
    f_val = Font(name="Arial", size=8, bold=True)

    gruppi_data = [
        ("GRUPPO ROSSO", squadra_rossa, data_r, avversario_rosso, in_casa_rosso, tecnico_r, tess_tec_r, dirigente_r, tess_dir_r, guardalinee_r, tess_gua_r, genera_distinta_r),
        ("GRUPPO BLU", squadra_blu, data_b, avversario_blu, in_casa_blu, tecnico_b, tess_tec_b, dirigente_b, tess_dir_b, guardalinee_b, tess_gua_b, genera_distinta_b)
    ]

    for nome_sheet, convocati, data_gara_g, avversario, in_casa, tecnico, tess_tec, dirigente, tess_dir, guardalinee, tess_gua, da_generare in gruppi_data:
        if not da_generare:
            continue

        ws = wb.create_sheet(title=nome_sheet)
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:C1")
        ws["A1"] = f"F.I.G.C. - Delegazione Provinciale/Distrettuale di"
        ws["A1"].font = f_sub
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("D1:R1")
        ws["D1"] = delegazione.upper()
        ws["D1"].font = f_title
        ws["D1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("S1:T1")
        ws["S1"] = "STAGIONE"
        ws["S1"].font = f_sub
        ws["S1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("U1:V1")
        ws["U1"] = stagione
        ws["U1"].font = f_val
        ws["U1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:C2")
        ws["A2"] = f"CATEG. {categoria}\nANNO {anno}"
        ws["A2"].font = f_sub
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws["D2"] = "Girone"
        ws["D2"].font = f_sub
        ws["D2"].alignment = Alignment(horizontal="center", vertical="center")

        sq_a = MIA_SQUADRA if in_casa else avversario
        sq_b = avversario if in_casa else MIA_SQUADRA

        ws.merge_cells("E2:K2")
        ws["E2"] = f"INCONTRO/CONFRONTO\n{sq_a} - {sq_b}"
        ws["E2"].font = f_sub
        ws["E2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("L2:M2")
        ws["L2"] = luogo
        ws["L2"].font = f_val
        ws["L2"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("N2:O2")
        ws["N2"] = f"DATA\n{data_gara_g}"
        ws["N2"].font = f_sub
        ws["N2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("P2:Q2")
        ws["P2"] = f"MODALITA' DI GIOCO\n{modalita}"
        ws["P2"].font = f_sub
        ws["P2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("R2:T2")
        ws["R2"] = f"MULTIPARTITA\nGARE  {num_gare}"
        ws["R2"].font = f_sub
        ws["R2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("U2:V2")
        ws["U2"] = "RISULTATO\n "
        ws["U2"].font = f_sub
        ws["U2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("A3:K3")
        ws["A3"] = 'S Q U A D R A "A"'
        ws["A3"].font = f_title
        ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("L3:V3")
        ws["L3"] = 'S Q U A D R A "B"'
        ws["L3"].font = f_title
        ws["L3"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A4:A5")
        ws["A4"] = "N°"
        ws["A4"].font = f_sub
        ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("B4:D4")
        ws["B4"] = "DATA DI NASCITA"
        ws["B4"].font = f_sub
        ws["B4"].alignment = Alignment(horizontal="center", vertical="center")

        ws["B5"] = "GG"
        ws["C5"] = "MM"
        ws["D5"] = "AA"
        for c in ["B5", "C5", "D5"]:
            ws[c].font = f_sub
            ws[c].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("E4:E5")
        ws["E4"] = "N° CARTELLINO"
        ws["E4"].font = f_sub
        ws["E4"].alignment = Alignment(horizontal="center", vertical="center")

        # SQUADRA A: COGNOME e NOME
        ws.merge_cells("F4:F5")
        ws["F4"] = "COGNOME"
        ws["F4"].font = f_sub
        ws["F4"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("G4:H5")
        ws["G4"] = "NOME"
        ws["G4"].font = f_sub
        ws["G4"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("I4:K4")
        ws["I4"] = "Presenza"
        ws["I4"].font = f_sub
        ws["I4"].alignment = Alignment(horizontal="center", vertical="center")

        ws["I5"] = "1T"
        ws["J5"] = "2T"
        ws["K5"] = "3T"
        for c in ["I5", "J5", "K5"]:
            ws[c].font = f_sub
            ws[c].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("L4:L5")
        ws["L4"] = "N°"
        ws["L4"].font = f_sub
        ws["L4"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("M4:O4")
        ws["M4"] = "DATA DI NASCITA"
        ws["M4"].font = f_sub
        ws["M4"].alignment = Alignment(horizontal="center", vertical="center")

        ws["M5"] = "GG"
        ws["N5"] = "MM"
        ws["O5"] = "AA"
        for c in ["M5", "N5", "O5"]:
            ws[c].font = f_sub
            ws[c].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("P4:P5")
        ws["P4"] = "N° CARTELLINO"
        ws["P4"].font = f_sub
        ws["P4"].alignment = Alignment(horizontal="center", vertical="center")

        # SQUADRA B: COGNOME e NOME
        ws.merge_cells("Q4:R5")
        ws["Q4"] = "COGNOME"
        ws["Q4"].font = f_sub
        ws["Q4"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("S4:S5")
        ws["S4"] = "NOME"
        ws["S4"].font = f_sub
        ws["S4"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("T4:V4")
        ws["T4"] = "Presenza"
        ws["T4"].font = f_sub
        ws["T4"].alignment = Alignment(horizontal="center", vertical="center")

        ws["T5"] = "1T"
        ws["U5"] = "2T"
        ws["V5"] = "3T"
        for c in ["T5", "U5", "V5"]:
            ws[c].font = f_sub
            ws[c].alignment = Alignment(horizontal="center", vertical="center")

        is_my_team_a = in_casa
        deve_compilare = (not in_casa)

        for r in range(1, 13):
            row_idx = r + 5
            ws.cell(row=row_idx, column=1, value=r).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row_idx, column=12, value=r).alignment = Alignment(horizontal="center", vertical="center")
            
            ws.merge_cells(start_row=row_idx, start_column=7, end_row=row_idx, end_column=8)
            ws.merge_cells(start_row=row_idx, start_column=17, end_row=row_idx, end_column=18)

            g_data = convocati[r-1] if (r <= len(convocati) and deve_compilare) else None

            if is_my_team_a and g_data:
                ws.cell(row=row_idx, column=2, value=g_data.get("g", "")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=3, value=g_data.get("m", "")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=4, value=str(g_data.get("a", ""))[-2:]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=5, value=g_data.get("Tesserino", "")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=6, value=g_data.get("Cognome", ""))
                ws.cell(row=row_idx, column=7, value=g_data.get("Nome", ""))

            if (not is_my_team_a) and g_data:
                ws.cell(row=row_idx, column=13, value=g_data.get("g", "")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=14, value=g_data.get("m", "")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=15, value=str(g_data.get("a", ""))[-2:]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=16, value=g_data.get("Tesserino", "")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=17, value=g_data.get("Cognome", ""))
                ws.cell(row=row_idx, column=19, value=g_data.get("Nome", ""))

        ws.merge_cells("A18:C18")
        ws["A18"] = "TECNICO A d B."
        ws.merge_cells("D18:G18")
        ws.merge_cells("H18:K18")
        ws["H18"] = "Tess. N°"

        ws.merge_cells("L18:N18")
        ws["L18"] = "TECNICO A.d.B."
        ws.merge_cells("O18:R18")
        ws.merge_cells("S18:V18")
        ws["S18"] = "Tess. N°"

        ws.merge_cells("A19:C19")
        ws["A19"] = "DIRIGENTE ACC."
        ws.merge_cells("D19:G19")
        ws.merge_cells("H19:K19")
        ws["H19"] = "Tess. N°"

        ws.merge_cells("L19:N19")
        ws["L19"] = "DIRIGENTE ACC."
        ws.merge_cells("O19:R19")
        ws.merge_cells("S19:V19")
        ws["S19"] = "Tess. N°"

        ws.merge_cells("A20:C20")
        ws["A20"] = "GUARDALINEE"
        ws.merge_cells("D20:G20")
        ws.merge_cells("H20:K20")
        ws["H20"] = "Tess. N°"

        ws.merge_cells("L20:N20")
        ws["L20"] = "GUARDALINEE"
        ws.merge_cells("O20:R20")
        ws.merge_cells("S20:V20")
        ws["S20"] = "Tess. N°"

        if deve_compilare:
            if is_my_team_a:
                ws["D18"] = tecnico
                ws["H18"] = f"Tess. N° {tess_tec}" if tess_tec else "Tess. N°"
                ws["D19"] = dirigente
                ws["H19"] = f"Tess. N° {tess_dir}" if tess_dir else "Tess. N°"
                ws["D20"] = guardalinee
                ws["H20"] = f"Tess. N° {tess_gua}" if tess_gua else "Tess. N°"
            else:
                ws["O18"] = tecnico
                ws["S18"] = f"Tess. N° {tess_tec}" if tess_tec else "Tess. N°"
                ws["O19"] = dirigente
                ws["S19"] = f"Tess. N° {tess_dir}" if tess_dir else "Tess. N°"
                ws["O20"] = guardalinee
                ws["S20"] = f"Tess. N° {tess_gua}" if tess_gua else "Tess. N°"

        ws.merge_cells("A21:K21")
        ws["A21"] = "SQUADRA 'A' — DA COMPILARE A CURA DEL DIRIGENTE DELLA SQUADRA 'B'"
        ws["A21"].font = f_sub
        ws["A21"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A22:K22")
        ws["A22"] = "SALUTO INIZIO E FINE GARA:  [  ] SUFFICIENTE   [  ] BUONO   [  ] OTTIMO"
        ws["A22"].font = f_small
        ws["A22"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A23:K23")
        ws["A23"] = "COMPORTAMENTO DIRIGENTI SQ. 'A': [  ] INSUFFICIENTE [  ] SUFFICIENTE [  ] BUONO [  ] OTTIMO"
        ws["A23"].font = f_small
        ws["A23"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A24:K24")
        ws["A24"] = "COMPORTAMENTO PUBBLICO SQ. 'A':  [  ] INSUFFICIENTE [  ] SUFFICIENTE [  ] BUONO [  ] OTTIMO"
        ws["A24"].font = f_small
        ws["A24"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("L21:V21")
        ws["L21"] = "SQUADRA 'B' — DA COMPILARE A CURA DEL DIRIGENTE DELLA SQUADRA 'A'"
        ws["L21"].font = f_sub
        ws["L21"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("L22:V22")
        ws["L22"] = "SALUTO INIZIO E FINE GARA:  [  ] SUFFICIENTE   [  ] BUONO   [  ] OTTIMO"
        ws["L22"].font = f_small
        ws["L22"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("L23:V23")
        ws["L23"] = "COMPORTAMENTO DIRIGENTI SQ. 'B': [  ] INSUFFICIENTE [  ] SUFFICIENTE [  ] BUONO [  ] OTTIMO"
        ws["L23"].font = f_small
        ws["L23"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("L24:V24")
        ws["L24"] = "COMPORTAMENTO PUBBLICO SQ. 'B':  [  ] INSUFFICIENTE [  ] SUFFICIENTE [  ] BUONO [  ] OTTIMO"
        ws["L24"].font = f_small
        ws["L24"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A25:K25")
        ws.merge_cells("L25:V25")

        ws.merge_cells("A26:K26")
        ws["A26"] = "FIRMA DIRIGENTE SQUADRA 'A': ____________________________________"
        ws["A26"].font = f_sub
        ws["A26"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("L26:V26")
        ws["L26"] = "FIRMA DIRIGENTE SQUADRA 'B': ____________________________________"
        ws["L26"].font = f_sub
        ws["L26"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A27:K27")
        ws["A27"] = "Da consegnare in busta chiusa o mezzo posta entro le 48 ore successive alla gara."
        ws["A27"].font = Font(name="Arial", size=7, italic=True)
        ws["A27"].alignment = Alignment(horizontal="left", vertical="center")

        for r_idx in range(1, 28):
            for c_idx in range(1, 23):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = border_thin

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

st.title("⚽ Convocazioni & Distinta Ufficiale FIGC AGAPE")

st.header("⚙️ Output Desiderato")
tipo_output = st.selectbox(
    "Cosa vuoi generare?",
    ["Tutto (Match Day + Convocazione + Distinta Excel)", "Solo Testi (Match Day + Convocazione)", "Solo Distinta Excel"]
)

is_solo_distinta = (tipo_output == "Solo Distinta Excel")

squadre_da_generare = st.radio(
    "Quale squadra/gruppo vuoi generare?",
    ["🔴🔵 Entrambe (Rosso e Blu)", "🔴 Solo Gruppo Rosso", "🔵 Solo Gruppo Blu"]
)

st.markdown("---")
st.header("⚙️ Parametri Generali Categoria")
delegazione = st.text_input("Delegazione FIGC", "CITTA' DI CASTELLO")
stagione = st.text_input("Stagione Sportiva", "26 / 27")
categoria = st.text_input("Categoria", "PRIMI CALCI")
anno = st.text_input("Anno Categoria", "2019")
modalita = st.text_input("Modalità di Gioco", "5 vs 5")
luogo = st.text_input("Luogo / Comune", "UMBERTIDE")
num_gare = st.text_input("Numero Gare Multipartita", "3")

st.markdown("---")

st.header("📋 Modalità di Selezione Convocati")
modalita_conv = st.radio(
    "Scegli come assegnare i giocatori alle squadre:",
    ["🎲 Casuale / Bilanciato (Automatico)", "☑️ Manuale (Seleziona Giocatori con Checkbox)"]
)

opzioni_giocatori = [f"{g['Cognome']} {g['Nome']} ({g['Livello']})" for g in ROSA]
mppa_giocatori = {f"{g['Cognome']} {g['Nome']} ({g['Livello']})": g for g in ROSA}

squadra_rossa_scelta = []
squadra_blu_scelta = []

if modalita_conv == "☑️ Manuale (Seleziona Giocatori con Checkbox)":
    if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔴 Solo Gruppo Rosso"]:
        st.subheader("🔴 Seleziona Giocatori per GRUPPO ROSSO")
        scelti_r = st.multiselect("Scegli i giocatori del Gruppo Rosso:", opzioni_giocatori, key="m_rosso")
        squadra_rossa_scelta = [mppa_giocatori[x] for x in scelti_r]
    else:
        scelti_r = []

    if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔵 Solo Gruppo Blu"]:
        st.subheader("🔵 Seleziona Giocatori per GRUPPO BLU")
        opzioni_blu = [x for x in opzioni_giocatori if x not in scelti_r]
        scelti_b = st.multiselect("Scegli i giocatori del Gruppo Blu:", opzioni_blu, key="m_blu")
        squadra_blu_scelta = [mppa_giocatori[x] for x in scelti_b]

gen_distinta_r = False
gen_distinta_b = False

if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔴 Solo Gruppo Rosso"]:
    st.markdown("---")
    st.header("🔴 PARAMETRI GARA GRUPPO ROSSO")
    gen_distinta_r = st.checkbox("Genera Distinta per Gruppo Rosso", value=True)
    data_r = st.text_input("📅 Data Gara Rosso (es. 24/10/2026)", "24/10/2026")
    avversario_r = st.text_input("Avversario Rosso", "SAN BIAGIO PROMANO")
    in_casa_r = st.checkbox("Gioca in Casa (Rosso)", value=False)
    tecnico_r = st.text_input("Mister Rosso", "Sig. TAVERNELLI ANDREA")
    tess_tec_r = st.text_input("Tess. Mister Rosso", "191991")
    dirigente_r = st.text_input("Dirigente Acc. Rosso", "Sig. OUAZINE IHAB")
    tess_dir_r = st.text_input("Tess. Dirigente Rosso", "")
    guardalinee_r = st.text_input("Guardalinee Rosso", "")
    tess_gua_r = st.text_input("Tess. Guardalinee Rosso", "")

    if not is_solo_distinta:
        ora_inizio_r = st.text_input("Ora Inizio Gara Rosso", "15:00")
        stadio_r = st.text_input("Stadio / Campo Rosso", "Stadio Comunale Umbertide")
        ritrovo_r = st.text_input("Ritrovo Rosso", "Campo Sportivo Umbertide")
        spogliatoio_r = st.text_input("Spogliatoio Rosso", "Spogliatoio 1")
        ora_conv_r = st.text_input("Ora Convocazione Rossa", "14:00")
    else:
        ora_inizio_r, stadio_r, ritrovo_r, spogliatoio_r, ora_conv_r = "", "", "", "", ""
else:
    data_r, ora_inizio_r, avversario_r, in_casa_r, stadio_r, ritrovo_r, spogliatoio_r, ora_conv_r, tecnico_r, tess_tec_r, dirigente_r, tess_dir_r, guardalinee_r, tess_gua_r = "", "", "", False, "", "", "", "", "", "", "", "", "", ""

if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔵 Solo Gruppo Blu"]:
    st.markdown("---")
    st.header("🔵 PARAMETRI GARA GRUPPO BLU")
    gen_distinta_b = st.checkbox("Genera Distinta per Gruppo Blu", value=True)
    data_b = st.text_input("📅 Data Gara Blu (es. 25/10/2026)", "25/10/2026")
    avversario_b = st.text_input("Avversario Blu", "CANNARA")
    in_casa_b = st.checkbox("Gioca in Casa (Blu)", value=False)
    tecnico_b = st.text_input("Mister Blu", "Sig. TAVERNELLI ANDREA")
    tess_tec_b = st.text_input("Tess. Mister Blu", "191991")
    dirigente_b = st.text_input("Dirigente Acc. Blu", "Sig. OUAZINE IHAB")
    tess_dir_b = st.text_input("Tess. Dirigente Blu", "")
    guardalinee_b = st.text_input("Guardalinee Blu", "")
    tess_gua_b = st.text_input("Tess. Guardalinee Blu", "")

    if not is_solo_distinta:
        ora_inizio_b = st.text_input("Ora Inizio Gara Blu", "10:30")
        stadio_b = st.text_input("Stadio / Campo Blu", "Stadio Spoletini Cannara")
        ritrovo_b = st.text_input("Ritrovo Blu", "Piazzale Parcheggio Umbertide")
        spogliatoio_b = st.text_input("Spogliatoio Blu", "Spogliatoio Ospiti")
        ora_conv_b = st.text_input("Ora Convocazione Blu", "09:30")
    else:
        ora_inizio_b, stadio_b, ritrovo_b, spogliatoio_b, ora_conv_b = "", "", "", "", ""
else:
    data_b, ora_inizio_b, avversario_b, in_casa_b, stadio_b, ritrovo_b, spogliatoio_b, ora_conv_b, tecnico_b, tess_tec_b, dirigente_b, tess_dir_b, guardalinee_b, tess_gua_b = "", "", "", False, "", "", "", "", "", "", "", "", "", ""

st.markdown("---")

if st.button("⚡ GENERATE CONVOCAZIONI / DISTINTA", type="primary", use_container_width=True):
    if not ROSA:
        st.error("Impossibile procedere: la rosa è vuota. Controlla il file giocatori.json.")
    else:
        if modalita_conv == "🎲 Casuale / Bilanciato (Automatico)":
            rossa, blu = genera_convocazione_random(ROSA)
        else:
            rossa, blu = squadra_rossa_scelta, squadra_blu_scelta

        casa_r = "AGAPE UMBERTIDE PIERANTONIO 🔴🔵" if in_casa_r else avversario_r
        trasf_r = avversario_r if in_casa_r else "AGAPE UMBERTIDE PIERANTONIO 🔴🔵"
        
        casa_b = "AGAPE UMBERTIDE PIERANTONIO 🔴🔵" if in_casa_b else avversario_b
        trasf_b = avversario_b if in_casa_b else "AGAPE UMBERTIDE PIERANTONIO 🔴🔵"

        if tipo_output in ["Tutto (Match Day + Convocazione + Distinta Excel)", "Solo Testi (Match Day + Convocazione)"]:
            matchday_blocks = []
            convocazioni_blocks = []

            if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔴 Solo Gruppo Rosso"]:
                matchday_blocks.append(f"""🔴GRUPPO ROSSO🔴\n📅Match Day – {data_r}\n⚽ {casa_r} vs {trasf_r}\nRb Inizio gara: ore {ora_inizio_r}\n📍 Stadio/Campo - {stadio_r}""")
                convocazioni_blocks.append(f"""📣 CONVOCAZIONE GARA\nGRUPPO ROSSO🔴\n📅 Data: {data_r}\n⏰ Convocazione: ore {ora_conv_r}\n📍 Ritrovo: {ritrovo_r}\n🚿 Spogliatoio : {spogliatoio_r}\n\n🕑 Inizio gara: ore {ora_inizio_r}\n🏟️ Campo: {stadio_r}\n\n⚔️ Gara: {casa_r} vs {trasf_r}\n\n👥 Convocati:\n{formatta_giocatori(rossa)}""")

            if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔵 Solo Gruppo Blu"]:
                matchday_blocks.append(f"""🔵GRUPPO BLU🔵\n📅 Match Day – {data_b}\n⚽ {casa_b} vs {trasf_b}\nRb Inizio gara: ore {ora_inizio_b}\n📍 Stadio/Campo - {stadio_b}""")
                convocazioni_blocks.append(f"""📣 CONVOCAZIONE GARA\nGRUPPO BLU 🔵\n📅 Data: {data_b}\n⏰ Convocazione: ore {ora_conv_b}\n📍 Ritrovo: {ritrovo_b}\n🚿 Spogliatoio : {spogliatoio_b}\n\n🕑 Inizio gara: ore {ora_inizio_b}\n🏟️ Campo: {stadio_b}\n\n⚔️ Gara: {casa_b} vs {trasf_b}\n\n👥 Convocati:\n{formatta_giocatori(blu)}""")

            matchday_txt = "\n\n".join(matchday_blocks)
            convocazione_txt = "📌 La convocazione seguirà nella giornata di giovedì:\n\n" + "\n\n".join(convocazioni_blocks) + "\n\n📲 *RISPONDERE AL SONDAGGIO ENTRO VENERDÌ ALLE ORE 13:00*"

            st.subheader("1. Testo per Match Day")
            st.code(matchday_txt, language=None)
            
            st.subheader("2. Testo Convocazioni (Giovedì)")
            st.code(convocazione_txt, language=None)

        if tipo_output in ["Tutto (Match Day + Convocazione + Distinta Excel)", "Solo Distinta Excel"]:
            gen_r = gen_distinta_r if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔴 Solo Gruppo Rosso"] else False
            gen_b = gen_distinta_b if squadre_da_generare in ["🔴🔵 Entrambe (Rosso e Blu)", "🔵 Solo Gruppo Blu"] else False

            excel_data = crea_distinta_excel(
                rossa, blu, gen_r, gen_b, delegazione, stagione, categoria, anno, modalita, luogo, num_gare,
                data_r, avversario_r, in_casa_r, tecnico_r, tess_tec_r, dirigente_r, tess_dir_r, guardalinee_r, tess_gua_r,
                data_b, avversario_b, in_casa_b, tecnico_b, tess_tec_b, dirigente_b, tess_dir_b, guardalinee_b, tess_gua_b
            )

            file_data_label = data_r if data_r else data_b
            st.subheader("3. Distinta Ufficiale FIGC (Excel)")
            st.download_button(
                label="📄 SCARICA DISTINTA UFFICIALE FIGC (EXCEL)",
                data=excel_data,
                file_name=f"Distinta_FIGC_{file_data_label.replace('/', '-')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
